import io
import json
import pytest
from fastapi.testclient import TestClient
import main
from models.candidate import Candidate, Question, Scorecard
from services.document_parser import DocumentParser
from services.candidate_repository import CandidateRepository
from services.interview_service import InterviewService, DEFAULT_BEHAVIORAL_QUESTIONS
from services.database_manager import DatabaseManager

pytestmark = pytest.mark.unit


class TestDocumentParser:
    def test_parse_txt(self):
        content = b"John Doe\nSoftware Engineer with 5 years experience in Python and FastAPI."
        result = DocumentParser.parse("resume.txt", content)
        assert "John Doe" in result
        assert "Python" in result

    def test_parse_docx(self):
        # Create a real in-memory docx
        import docx
        doc = docx.Document()
        doc.add_paragraph("Jane Smith")
        doc.add_paragraph("Senior AI Engineer with experience in PyTorch and LLMs.")
        bio = io.BytesIO()
        doc.save(bio)
        data = bio.getvalue()

        result = DocumentParser.parse("resume.docx", data)
        assert "Jane Smith" in result
        assert "Senior AI Engineer" in result

    def test_parse_fallback(self):
        content = b"Candidate Profile plain bytes"
        result = DocumentParser.parse("unknown_ext.xyz", content)
        assert "Candidate Profile" in result


class TestCandidateRepository:
    def test_save_and_retrieve_fallback(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        # Force offline db to test file fallback
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates.json")

        cand = Candidate(
            id="cand_test_001",
            name="Alice Walker",
            phone="+1234567890",
            position="Staff ML Engineer",
            job_description="Lead ML infra",
            resume_filename="resume.pdf",
            resume_text="Alice's Resume content",
            questions=[
                Question(id="q1", category="behavioral", text="Tell me about a crisis", order=1),
                Question(id="q2", category="technical", text="Explain Transformers", order=2),
            ],
        )

        assert repo.save(cand) is True
        retrieved = repo.get_by_id("cand_test_001")
        assert retrieved is not None
        assert retrieved.name == "Alice Walker"
        assert retrieved.position == "Staff ML Engineer"
        assert len(retrieved.questions) == 2
        assert retrieved.questions[0].text == "Tell me about a crisis"

        # Test responses query tracked by candidate_id
        responses = repo.get_responses("cand_test_001")
        assert len(responses) == 2
        assert responses[0]["candidate_id"] == "cand_test_001"
        assert responses[0]["candidate_name"] == "Alice Walker"
        assert responses[0]["candidate_position"] == "Staff ML Engineer"

        # Test phone lookup
        by_phone = repo.get_by_phone("+1234567890")
        assert by_phone is not None
        assert by_phone.id == "cand_test_001"

        all_cands = repo.get_all()
        assert len(all_cands) == 1

        assert repo.delete("cand_test_001") is True
        assert repo.get_by_id("cand_test_001") is None

    def test_question_deletion_and_list_detail_consistency(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_sync.json")

        cand = Candidate(
            id="cand_sync_001",
            name="Bob Sync",
            phone="+919999911111",
            position="Backend Dev",
            questions=[
                Question(id="q1", category="technical", text="Question 1", order=1),
                Question(id="q2", category="technical", text="Question 2", order=2),
                Question(id="q3", category="behavioral", text="Question 3", order=3),
            ],
        )
        repo.save(cand)

        # Verify initial 3 questions
        loaded = repo.get_by_id("cand_sync_001")
        assert len(loaded.questions) == 3

        # Delete q2 and q3, keeping only q1
        cand.questions = [
            Question(id="q1", category="technical", text="Question 1 updated", order=1),
        ]
        repo.save(cand)

        # Detail view must have ONLY 1 question (no resurrection of q2, q3)
        detail = repo.get_by_id("cand_sync_001")
        assert len(detail.questions) == 1
        assert detail.questions[0].id == "q1"
        assert detail.questions[0].text == "Question 1 updated"

        # List view (get_all) must also have ONLY 1 question
        all_cands = repo.get_all()
        assert len(all_cands) == 1
        assert len(all_cands[0].questions) == 1
        assert all_cands[0].questions[0].id == "q1"

        # get_responses must also return ONLY 1 question
        responses = repo.get_responses("cand_sync_001")
        assert len(responses) == 1
        assert responses[0]["question_id"] == "q1"

    def test_underscore_ids_no_collision(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_underscore.json")

        # Two candidates whose concatenated IDs could collide if naive split was used
        cand1 = Candidate(
            id="cand_1_2",
            name="Candidate A",
            phone="111",
            position="Engineer",
            questions=[Question(id="3", category="technical", text="Q from Cand 1", order=1)],
        )
        cand2 = Candidate(
            id="cand_1",
            name="Candidate B",
            phone="222",
            position="Engineer",
            questions=[Question(id="2_3", category="technical", text="Q from Cand 2", order=1)],
        )
        repo.save(cand1)
        repo.save(cand2)

        r1 = repo.get_by_id("cand_1_2")
        r2 = repo.get_by_id("cand_1")
        assert len(r1.questions) == 1
        assert r1.questions[0].text == "Q from Cand 1"
        assert len(r2.questions) == 1
        assert r2.questions[0].text == "Q from Cand 2"

    def test_pagination_and_search_fallback(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_page.json")

        for i in range(5):
            repo.save(Candidate(
                id=f"c_{i}",
                name=f"Candidate {i}",
                phone=f"1000{i}",
                position="Dev" if i % 2 == 0 else "Designer",
                job_description="JD",
                status="ready" if i < 3 else "evaluated",
            ))

        assert repo.get_count() == 5
        assert repo.get_count(status="ready") == 3
        assert repo.get_count(search="Candidate 1") == 1

        # Page 1 (limit 2)
        p1 = repo.get_all(limit=2, offset=0)
        assert len(p1) == 2

        # Filter by status
        ready = repo.get_all(status="ready")
        assert len(ready) == 3

    def test_db_error_consistent_fallback_queries(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        # Mock connection where cursor operations raise an exception
        class FailingCursor:
            def __enter__(self):
                return self
            def __exit__(self, exc_type, exc_val, exc_tb):
                pass
            def execute(self, *args, **kwargs):
                raise Exception("Simulated PostgreSQL connection failure")
            def fetchall(self):
                raise Exception("Simulated PostgreSQL connection failure")
            def fetchone(self):
                raise Exception("Simulated PostgreSQL connection failure")

        class FailingConn:
            def cursor(self, *args, **kwargs):
                return FailingCursor()
            def close(self):
                pass

        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_err.json")

        # Seed fallback store with 4 candidates
        for i in range(4):
            repo.save(Candidate(
                id=f"cand_err_{i}",
                name=f"Dev {i}",
                phone=f"+91100{i}",
                position="Backend Engineer",
                status="ready" if i % 2 == 0 else "evaluated",
                questions=[Question(id=f"q_{i}", text=f"Question {i}", category="technical", order=1)],
            ))

        # Now simulate DB connection that throws errors on queries
        monkeypatch.setattr(db_mgr, "get_connection", lambda: FailingConn())

        # 1. get_count() on DB error must NOT return 0; must return filtered fallback count
        assert repo.get_count(status="ready") == 2
        assert repo.get_count(status="evaluated") == 2
        assert repo.get_count(search="Dev 1") == 1

        # 2. get_all() on DB error must apply filters and pagination, not just [:limit]
        res_ready = repo.get_all(limit=1, offset=0, status="ready")
        assert len(res_ready) == 1
        assert res_ready[0].status == "ready"

        res_paged = repo.get_all(limit=2, offset=1)
        assert len(res_paged) == 2

        # 3. get_responses() on DB error must return candidate's questions, not empty list []
        responses = repo.get_responses("cand_err_0")
        assert len(responses) == 1
        assert responses[0]["question_id"] == "q_0"

    def test_delete_purges_fallback_file(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        # Simulate successful DB connection that reports 1 row deleted
        class SuccessfulCursor:
            def __init__(self):
                self.rowcount = 1
            def __enter__(self):
                return self
            def __exit__(self, *args):
                pass
            def execute(self, query, *args, **kwargs):
                if "SELECT id FROM candidates" in query:
                    pass
                elif "DELETE FROM candidates" in query:
                    self.rowcount = 1
            def fetchone(self):
                return ("c_del_1",)
            def fetchall(self):
                return []

        class SuccessfulConn:
            def cursor(self, *args, **kwargs):
                return SuccessfulCursor()
            def close(self):
                pass

        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_del.json")

        # First save candidate to fallback
        cand = Candidate(id="c_del_1", name="Del Me", phone="123", position="QA")
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        repo.save(cand)
        assert "c_del_1" in repo._load_fallback()

        # Now simulate successful DB delete; fallback copy must also be removed
        monkeypatch.setattr(db_mgr, "get_connection", lambda: SuccessfulConn())
        assert repo.delete("c_del_1") is True
        assert "c_del_1" not in repo._load_fallback()

    def test_delete_nonexistent_returns_false(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_del_none.json")

        # 1. In offline mode: nonexistent candidate returns False
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        assert repo.delete("cand_ghost_999") is False

        # 2. In DB mode: when row doesn't exist in DB, delete returns False
        class NonexistentCursor:
            def __init__(self):
                self.rowcount = 0
            def __enter__(self):
                return self
            def __exit__(self, *args):
                pass
            def execute(self, *args, **kwargs):
                self.rowcount = 0
            def fetchone(self):
                return None
            def fetchall(self):
                return []

        class NonexistentConn:
            def cursor(self, *args, **kwargs):
                return NonexistentCursor()
            def close(self):
                pass

        monkeypatch.setattr(db_mgr, "get_connection", lambda: NonexistentConn())
        assert repo.delete("cand_ghost_999") is False

    def test_delete_purges_recordings_and_calls(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        recordings_dir = tmp_path / "recordings"
        recordings_dir.mkdir()
        dummy_wav = recordings_dir / "recording_call_99.wav"
        dummy_wav.write_bytes(b"RIFF dummy wav audio data")

        from core.config import settings
        monkeypatch.setattr(settings, "RECORDINGS_DIR", recordings_dir)

        deleted_queries = []
        class CallCursor:
            def __init__(self):
                self.rowcount = 1
            def __enter__(self):
                return self
            def __exit__(self, *args):
                pass
            def execute(self, query, *args, **kwargs):
                deleted_queries.append(query)
                self.rowcount = 1
            def fetchone(self):
                return ("c_with_call",)
            def fetchall(self):
                return [("recording_call_99.wav",)]

        class CallConn:
            def cursor(self, *args, **kwargs):
                return CallCursor()
            def close(self):
                pass

        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_media.json")
        monkeypatch.setattr(db_mgr, "get_connection", lambda: CallConn())

        # Verify wav exists before deletion
        assert dummy_wav.exists()

        assert repo.delete("c_with_call", purge_linked_calls=True, purge_recordings=True) is True

        # Verify physical recording WAV was deleted from disk
        assert not dummy_wav.exists()

        # Verify calls and call_logs deletion queries were executed
        assert any("DELETE FROM calls" in q for q in deleted_queries)
        assert any("DELETE FROM call_logs" in q for q in deleted_queries)


    def test_offline_save_dirty_tracking_and_replay(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_replay.json")

        # 1. Save while offline
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        cand = Candidate(
            id="c_rep_1",
            name="Replay Candidate",
            phone="987",
            position="SRE",
            questions=[Question(id="q1", text="SRE question", category="technical")],
        )
        repo.save(cand)

        # Verify dirty tracking and status visibility
        assert repo.last_write_storage == "fallback"
        assert repo.get_pending_sync_count() == 1
        status = repo.get_storage_status()
        assert status["mode"] == "fallback"
        assert status["pending_sync_count"] == 1
        assert status["fallback_records_count"] == 1

        # 2. Simulate DB recovery and replay
        saved_to_db = []
        monkeypatch.setattr(repo, "_save_to_postgres", lambda c, conn: saved_to_db.append(c.id))

        class MockConn:
            def close(self):
                pass

        monkeypatch.setattr(db_mgr, "get_connection", lambda: MockConn())
        sync_result = repo.sync_fallback_to_db()

        assert sync_result["synced_count"] == 1
        assert sync_result["pending_count"] == 0
        assert "c_rep_1" in saved_to_db
        assert repo.get_pending_sync_count() == 0

    def test_concurrent_fallback_writes(self, tmp_path, monkeypatch):
        import concurrent.futures
        db_mgr = DatabaseManager()
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_concurrent.json")

        num_threads = 10
        candidates = [
            Candidate(id=f"conc_{i}", name=f"Concurrent {i}", phone=f"555{i}", position="Dev")
            for i in range(num_threads)
        ]

        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
            futures = [executor.submit(repo.save, c) for c in candidates]
            for f in concurrent.futures.as_completed(futures):
                assert f.result() is True

        data = repo._load_fallback()
        assert len(data) == num_threads
        for i in range(num_threads):
            assert f"conc_{i}" in data

    def test_phone_normalization_and_exact_matching(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        monkeypatch.setattr(db_mgr, "get_connection", lambda: None)
        repo = CandidateRepository(db_mgr)
        repo.fallback_file = str(tmp_path / "candidates_phone.json")

        cand_us = Candidate(
            id="cand_us_01",
            name="Alice US",
            phone="+1 (415) 555-2671",
            position="Engineer",
        )
        cand_in = Candidate(
            id="cand_in_01",
            name="Raj India",
            phone="+91 98765-43210",
            position="Lead",
        )
        repo.save(cand_us)
        repo.save(cand_in)

        # 1. Exact match with variations
        assert repo.get_by_phone("+14155552671").id == "cand_us_01"
        assert repo.get_by_phone("4155552671").id == "cand_us_01"
        assert repo.get_by_phone("(415) 555-2671").id == "cand_us_01"

        assert repo.get_by_phone("+919876543210").id == "cand_in_01"
        assert repo.get_by_phone("9876543210").id == "cand_in_01"
        assert repo.get_by_phone("+91 98765-43210").id == "cand_in_01"

        # 2. Rejection of empty, whitespace, and short inputs
        assert repo.get_by_phone("") is None
        assert repo.get_by_phone("   ") is None
        assert repo.get_by_phone("+") is None
        assert repo.get_by_phone("123") is None
        assert repo.get_by_phone("0") is None

        # 3. Elimination of false-positive substring matches
        assert repo.get_by_phone("98765") is None
        assert repo.get_by_phone("43210") is None
        assert repo.get_by_phone("5552671") is None
        assert repo.get_by_phone("+14155552672") is None  # 1 digit off

    def test_phone_btree_index_query_not_ilike(self, tmp_path, monkeypatch):
        db_mgr = DatabaseManager()
        executed_queries = []

        class QueryTrackingCursor:
            def __enter__(self):
                return self
            def __exit__(self, *args):
                pass
            def execute(self, query, params=None):
                executed_queries.append((query, params))
            def fetchone(self):
                return ("cand_btree_01",)

        class QueryTrackingConn:
            def cursor(self, *args, **kwargs):
                return QueryTrackingCursor()
            def close(self):
                pass

        repo = CandidateRepository(db_mgr)
        monkeypatch.setattr(db_mgr, "get_connection", lambda: QueryTrackingConn())
        monkeypatch.setattr(repo, "get_by_id", lambda cid: Candidate(id=cid, name="BTree Cand", phone="+14155552671", position="Dev"))

        cand = repo.get_by_phone("+1 (415) 555-2671")
        assert cand is not None
        assert cand.id == "cand_btree_01"

        # Verify query uses exact equality with ANY(%s) and NO wildcard ILIKE / LIKE
        assert len(executed_queries) == 1
        sql, params = executed_queries[0]
        assert "WHERE phone = ANY(%s)" in sql
        assert "ILIKE" not in sql
        assert "LIKE" not in sql
        # Check params contains exact variants without any '%' wildcards
        variants = params[0]
        assert not any("%" in v for v in variants)
        assert "+14155552671" in variants
        assert "4155552671" in variants


class TestInterviewService:
    @pytest.mark.asyncio
    async def test_default_behavioral_questions(self):
        service = InterviewService()
        questions = await service.generate_questions(
            candidate_name="Bob",
            position="DevOps Engineer",
            job_description="Kubernetes, Terraform, AWS",
            resume_text="Bob's background in AWS and CI/CD",
        )
        assert len(questions) >= 5
        beh_questions = [q for q in questions if q.category == "behavioral"]
        assert len(beh_questions) == 5
        assert "deadline" in beh_questions[0].text.lower() or "crisis" in beh_questions[0].text.lower()

    @pytest.mark.asyncio
    async def test_custom_behavioral_questions(self):
        service = InterviewService()
        custom = [f"Custom Q{i}" for i in range(1, 6)]
        questions = await service.generate_questions(
            candidate_name="Carol",
            position="Frontend Lead",
            job_description="React, TypeScript, CSS",
            resume_text="Carol's React resume",
            custom_behavioral=custom,
        )
        beh_questions = [q for q in questions if q.category == "behavioral"]
        assert len(beh_questions) == 5
        assert beh_questions[0].text == "Custom Q1"
        assert beh_questions[4].text == "Custom Q5"

    @pytest.mark.asyncio
    async def test_evaluate_interview_empty_transcript(self):
        service = InterviewService()
        cand = Candidate(
            id="c1",
            name="Dan",
            phone="123",
            position="QA",
            job_description="Automation",
            resume_filename="r.pdf",
            resume_text="Dan resume",
        )
        scorecard = await service.evaluate_interview(cand, [])
        assert scorecard.recommendation == "Do Not Hire"
        assert scorecard.overall_score == 0


class TestCandidateRoutes:
    def test_behavioral_defaults_endpoint(self):
        client = TestClient(main.app)
        resp = client.get("/api/candidates/behavioral-defaults")
        assert resp.status_code == 200
        body = resp.json()
        assert "questions" in body
        assert len(body["questions"]) == 5

    def test_candidate_crud_cycle(self, monkeypatch, tmp_path):
        from services.candidate_repository import candidate_repo
        candidate_repo.fallback_file = str(tmp_path / "candidates_test.json")
        monkeypatch.setattr(candidate_repo.db_manager, "get_connection", lambda: None)

        client = TestClient(main.app)

        # 1. Create candidate via multipart form
        file_payload = ("resume.txt", b"Jane Doe\nFull Stack Developer with Node and React", "text/plain")
        form_data = {
            "name": "Jane Doe",
            "phone": "+91 9988776655",
            "position": "Full Stack Developer",
            "job_description": "We are seeking a Full Stack Developer experienced with Node, React, and PostgreSQL.",
        }
        res = client.post("/api/candidates", data=form_data, files={"resume": file_payload})
        assert res.status_code == 200
        cand_data = res.json()
        cand_id = cand_data["id"]
        assert cand_data["name"] == "Jane Doe"
        assert cand_data["position"] == "Full Stack Developer"
        assert len(cand_data["questions"]) >= 5

        # 2. Get all candidates
        list_res = client.get("/api/candidates")
        assert list_res.status_code == 200
        items = list_res.json()
        assert any(c["id"] == cand_id for c in items)

        # 3. Get candidate by id
        get_res = client.get(f"/api/candidates/{cand_id}")
        assert get_res.status_code == 200
        assert get_res.json()["name"] == "Jane Doe"

        # 4. Update questions
        updated_questions = [
            {"id": "q1", "category": "behavioral", "text": "Edited behavioral question", "order": 1, "completed": False}
        ]
        put_res = client.put(f"/api/candidates/{cand_id}/questions", json={"questions": updated_questions})
        assert put_res.status_code == 200
        assert len(put_res.json()["questions"]) == 1
        assert put_res.json()["questions"][0]["text"] == "Edited behavioral question"

        # 5. Get tracked candidate responses by candidate_id
        resp_res = client.get(f"/api/candidates/{cand_id}/responses")
        assert resp_res.status_code == 200
        assert resp_res.json()["candidate_id"] == cand_id
        assert len(resp_res.json()["responses"]) == 1
        assert resp_res.json()["responses"][0]["candidate_id"] == cand_id

        # 6. Get candidate calls by candidate_id
        calls_res = client.get(f"/api/candidates/{cand_id}/calls")
        assert calls_res.status_code == 200
        assert calls_res.json()["candidate_id"] == cand_id

        # 7. Evaluate candidate
        eval_res = client.post(f"/api/candidates/{cand_id}/evaluate")
        assert eval_res.status_code == 200
        assert eval_res.json()["status"] == "evaluated"
        assert eval_res.json()["scorecard"] is not None

        # 8. Delete candidate
        del_res = client.delete(f"/api/candidates/{cand_id}")
        assert del_res.status_code == 200

        # Verify deleted
        assert client.get(f"/api/candidates/{cand_id}").status_code == 404


class TestInterviewControlMarkers:
    def test_end_call_marker_sets_flag_and_is_stripped(self, session_state):
        from routes import voice
        out = voice.extract_control_markers("Thank you for joining today's interview. [END_CALL]", session_state)
        assert session_state["pending_end_call"] is True
        assert "END_CALL" not in out
        assert out == "Thank you for joining today's interview."

    def test_add_transcript_turn(self, session_state):
        from routes import voice
        voice.add_transcript(session_state, "caller", "I have 5 years experience with Python", "en-IN")
        assert len(session_state["transcript"]) == 1
        turn = session_state["transcript"][0]
        assert turn["role"] == "caller"
        assert turn["text"] == "I have 5 years experience with Python"
        assert turn["lang"] == "en-IN"


class TestSettingsUsageInServices:
    def test_interview_service_uses_settings(self):
        from services.interview_service import (
            SARVAM_LLM_BASE_URL,
            CEREBRAS_LLM_BASE_URL,
            GEMINI_LLM_BASE_URL,
            TOGETHER_LLM_BASE_URL,
            FALLBACK_MODEL,
        )
        from core.config import settings
        assert SARVAM_LLM_BASE_URL == settings.SARVAM_LLM_BASE_URL
        assert CEREBRAS_LLM_BASE_URL == settings.CEREBRAS_LLM_BASE_URL
        assert GEMINI_LLM_BASE_URL == settings.GEMINI_LLM_BASE_URL
        assert TOGETHER_LLM_BASE_URL == settings.TOGETHER_LLM_BASE_URL
        assert FALLBACK_MODEL == settings.FALLBACK_MODEL

    def test_voice_llm_uses_settings(self):
        from services.voice.llm import (
            SARVAM_LLM_BASE_URL,
            CEREBRAS_LLM_BASE_URL,
            GEMINI_LLM_BASE_URL,
            TOGETHER_LLM_BASE_URL,
            FALLBACK_MODEL,
        )
        from core.config import settings
        assert SARVAM_LLM_BASE_URL == settings.SARVAM_LLM_BASE_URL
        assert CEREBRAS_LLM_BASE_URL == settings.CEREBRAS_LLM_BASE_URL
        assert GEMINI_LLM_BASE_URL == settings.GEMINI_LLM_BASE_URL
        assert TOGETHER_LLM_BASE_URL == settings.TOGETHER_LLM_BASE_URL
        assert FALLBACK_MODEL == settings.FALLBACK_MODEL

    def test_voice_tts_and_stt_use_settings(self):
        from services.voice.tts import SARVAM_TTS_URL, CARTESIA_TTS_URL, CARTESIA_VERSION
        from services.voice.stt import SARVAM_STT_URL
        from services.translation_service import TranslationService
        from core.config import settings
        assert SARVAM_TTS_URL == settings.SARVAM_TTS_URL
        assert CARTESIA_TTS_URL == settings.CARTESIA_TTS_URL
        assert CARTESIA_VERSION == settings.CARTESIA_VERSION
        assert SARVAM_STT_URL == settings.SARVAM_STT_URL
        assert TranslationService.SARVAM_TRANSLATE_URL == settings.SARVAM_TRANSLATE_URL
