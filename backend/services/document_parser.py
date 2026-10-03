import io
import logging
from typing import BinaryIO
import pypdf
import docx

logger = logging.getLogger("services.document_parser")


class DocumentParser:
    """Extracts plain text content from uploaded candidate resumes (PDF, DOCX, TXT)."""

    @staticmethod
    def parse_pdf(stream: BinaryIO | bytes) -> str:
        """Extract text from a PDF file stream or bytes."""
        try:
            if isinstance(stream, bytes):
                stream = io.BytesIO(stream)
            reader = pypdf.PdfReader(stream)
            text_parts = []
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
            return "\n\n".join(text_parts).strip()
        except Exception as e:
            logger.error(f"[DOC PARSER] PDF extraction failed: {e}")
            raise ValueError(f"Failed to extract text from PDF: {e}")

    @staticmethod
    def parse_docx(stream: BinaryIO | bytes) -> str:
        """Extract text from a DOCX file stream or bytes."""
        try:
            if isinstance(stream, bytes):
                stream = io.BytesIO(stream)
            doc = docx.Document(stream)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs).strip()
        except Exception as e:
            logger.error(f"[DOC PARSER] DOCX extraction failed: {e}")
            raise ValueError(f"Failed to extract text from DOCX: {e}")

    @classmethod
    def parse(cls, filename: str, content: bytes) -> str:
        """Inspects extension and dispatches to appropriate parser."""
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            return cls.parse_pdf(content)
        elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
            return cls.parse_docx(content)
        elif lower_name.endswith(".txt"):
            try:
                return content.decode("utf-8").strip()
            except UnicodeDecodeError:
                return content.decode("latin-1", errors="replace").strip()
        else:
            # Fallback: attempt PDF first, then UTF-8
            try:
                return cls.parse_pdf(content)
            except Exception:
                return content.decode("utf-8", errors="ignore").strip()


document_parser = DocumentParser()
